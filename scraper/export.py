"""
取得結果を Excel / CSV / Text / JSON にエクスポートする
"""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# 1行あたりの最大文字数（Excelのセル制限等を考慮）
TEXT_SNIPPET_LEN = 8000


def _row_dict(item: dict) -> dict:
    """1件の取得結果を出力用の辞書に正規化"""
    text = (item.get("text") or "")[:TEXT_SNIPPET_LEN]
    return {
        "url": item.get("url", ""),
        "title": item.get("title") or "",
        "text": text,
        "status": item.get("status", "ok"),
        "fetched_at": item.get("fetched_at", ""),
    }


def export_excel(items: list[dict]) -> bytes:
    """取得結果をExcelバイナリで返す"""
    wb = Workbook()
    ws = wb.active
    ws.title = "取得結果"
    headers = ["URL", "タイトル", "本文", "ステータス", "取得日時"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for item in items:
        row = _row_dict(item)
        ws.append([
            row["url"],
            row["title"],
            row["text"],
            row["status"],
            row["fetched_at"],
        ])
    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.column_dimensions["C"].width = 50
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_csv(items: list[dict]) -> str:
    """取得結果をCSV文字列で返す"""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["URL", "タイトル", "本文", "ステータス", "取得日時"])
    for item in items:
        row = _row_dict(item)
        writer.writerow([
            row["url"],
            row["title"],
            row["text"],
            row["status"],
            row["fetched_at"],
        ])
    return buf.getvalue()


def export_text(items: list[dict]) -> str:
    """取得結果をプレーンテキストで返す（1件ずつ区切り）"""
    lines = []
    for i, item in enumerate(items, 1):
        row = _row_dict(item)
        lines.append(f"========== {i} ==========")
        lines.append(f"URL: {row['url']}")
        lines.append(f"タイトル: {row['title']}")
        lines.append(f"ステータス: {row['status']}")
        lines.append(f"取得日時: {row['fetched_at']}")
        lines.append("")
        lines.append(row["text"] or "(本文なし)")
        lines.append("")
    return "\n".join(lines)


def export_json(items: list[dict]) -> str:
    """取得結果をJSON文字列で返す"""
    out = [_row_dict(item) for item in items]
    return json.dumps(out, ensure_ascii=False, indent=2)
